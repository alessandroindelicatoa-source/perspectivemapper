from __future__ import annotations

import hashlib
import hmac
import io
import json
import os
import re
import sys
import warnings
import zipfile
import html as html_lib
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

from analysis_core import (
    DEFAULT_FRAMES, EMBEDDING_MODEL, SENTIMENT_MODEL,
    aggregate_embeddings, build_chunks, clean_for_lexical, collocations_pmi,
    document_statistics, excel_bytes, frame_indicators, hierarchy_from_embeddings,
    informative_log_odds, json_safe, kwic, numeric_group_test, safe_language,
    semantic_clustering, tokenize, top_tfidf_terms, topic_models,
    aggregate_entities, actor_topic_edges, bootstrap_difference_ci, bootstrap_mean_ci,
    robust_time_trend, supervised_embedding_classifier, temporal_bootstrap_summary,
    filter_documents_by_keywords, grouped_numeric_summary, parse_keyword_terms,
    aggregate_unit_topics, clean_extracted_text, concept_projection, dictionary_term_counts,
    find_repeated_page_lines, lemmatizer_available, parse_concept_anchors,
    topic_coherence_npmi, topic_diversity, topic_k_diagnostics,
)
from i18n import UI_LANGUAGES, tr

st.set_page_config(page_title="PerspectiveMapper v3.3", page_icon="🧭", layout="wide", initial_sidebar_state="expanded")

ROOT = Path(__file__).parent
NER_MODEL = "Davlan/xlm-roberta-base-ner-hrl"


# ---------- Styling ----------
st.markdown("""
<style>
.block-container {padding-top: 1.35rem; padding-bottom: 3rem; max-width: 1500px;}
[data-testid="stMetricValue"] {font-size: 1.7rem;}
.pm-note {padding: .75rem 1rem; border: 1px solid rgba(128,128,128,.25); border-radius: .6rem; margin: .5rem 0 1rem 0;}
.small-muted {opacity:.72; font-size:.9rem;}
</style>
""", unsafe_allow_html=True)


def password_matches(stored: str, supplied: str) -> bool:
    """Constant-time comparison; supports plain text or ``sha256:<hex>`` stored values."""
    if stored.startswith("sha256:"):
        supplied = "sha256:" + hashlib.sha256(supplied.encode("utf-8")).hexdigest()
        return hmac.compare_digest(stored.lower(), supplied.lower())
    return hmac.compare_digest(stored.encode("utf-8"), supplied.encode("utf-8"))


def optional_gate(lang: str) -> None:
    """Authenticate only when [passwords] exists in Streamlit secrets."""
    try:
        configured = "passwords" in st.secrets and len(st.secrets["passwords"]) > 0
    except Exception:
        configured = False
    if not configured:
        return
    if st.session_state.get("authenticated"):
        return
    st.title(f"🔐 PerspectiveMapper — {tr(lang, 'access')}")
    u = st.text_input(tr(lang, "username"))
    p = st.text_input(tr(lang, "password"), type="password")
    if st.button(tr(lang, "login"), type="primary"):
        try:
            if u in st.secrets["passwords"] and password_matches(str(st.secrets["passwords"][u]), str(p)):
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error(tr(lang, "bad_login"))
        except Exception:
            st.error(tr(lang, "bad_login"))
    st.stop()


@st.cache_resource(show_spinner=False)
def get_embedder():
    from sentence_transformers import SentenceTransformer
    return SentenceTransformer(EMBEDDING_MODEL)


@st.cache_resource(show_spinner=False)
def get_sentiment_pipeline():
    from transformers import pipeline
    return pipeline("text-classification", model=SENTIMENT_MODEL, tokenizer=SENTIMENT_MODEL, top_k=None, device=-1)


@st.cache_resource(show_spinner=False)
def get_ner_pipeline():
    from transformers import pipeline
    return pipeline("token-classification", model=NER_MODEL, tokenizer=NER_MODEL, aggregation_strategy="simple", device=-1)


def parse_sentiment_output(raw) -> Dict[str, float]:
    """Normalize transformers pipeline output across top_k/return_all_scores versions."""
    if isinstance(raw, list) and len(raw) == 1 and isinstance(raw[0], list):
        raw = raw[0]
    if isinstance(raw, dict):
        raw = [raw]
    probs = {"positive": np.nan, "neutral": np.nan, "negative": np.nan}
    for item in raw or []:
        label = str(item.get("label", "")).lower()
        score = float(item.get("score", np.nan))
        if "pos" in label or label == "label_0": probs["positive"] = score
        elif "neu" in label or label == "label_1": probs["neutral"] = score
        elif "neg" in label or label == "label_2": probs["negative"] = score
    return probs


@st.cache_data(show_spinner=False, max_entries=4)
def embed_texts(texts: Tuple[str, ...], _progress=None) -> np.ndarray:
    """Passage embeddings, cached on the exact passage texts so that changing topic or
    dictionary settings does not recompute them."""
    embedder = get_embedder()
    out = []
    step = 256
    for start in range(0, len(texts), step):
        if _progress is not None:
            _progress(start / max(len(texts), 1), f"{start:,}/{len(texts):,}")
        out.append(embedder.encode(list(texts[start:start + step]), batch_size=32,
                                   show_progress_bar=False, normalize_embeddings=True))
    return np.vstack(out) if out else np.zeros((0, 384))


@st.cache_data(show_spinner=False, max_entries=4)
def sentiment_raw(texts: Tuple[str, ...], _progress=None) -> list:
    pipe = get_sentiment_pipeline()
    outs = []
    batch = 16
    for start in range(0, len(texts), batch):
        if _progress is not None and start % (batch * 8) == 0:
            _progress(start / max(len(texts), 1), f"{start:,}/{len(texts):,}")
        outs.extend(pipe(list(texts[start:start + batch]), truncation=True, max_length=384, batch_size=batch))
    return outs


def sentiment_chunks(chunks: pd.DataFrame, progress=None) -> Tuple[pd.DataFrame, pd.DataFrame]:
    texts = chunks["chunk_text"].astype(str).tolist()
    raw_all = sentiment_raw(tuple(texts), progress)
    rows = []
    batch = len(texts) or 1
    for start in range(0, len(texts), batch):
        outs = raw_all[start:start+batch]
        for j, raw in enumerate(outs):
            probs = parse_sentiment_output(raw)
            score = probs["positive"] - probs["negative"] if np.isfinite(probs["positive"]) and np.isfinite(probs["negative"]) else np.nan
            label = max(probs, key=lambda k: probs[k] if np.isfinite(probs[k]) else -1)
            rows.append({
                "chunk_id": chunks.iloc[start+j]["chunk_id"],
                "doc_id": chunks.iloc[start+j]["doc_id"],
                "sentiment": label,
                "sentiment_score": score,
                "p_positive": probs["positive"],
                "p_neutral": probs["neutral"],
                "p_negative": probs["negative"],
            })
    cdf = pd.DataFrame(rows)
    cdf["chunk_words"] = chunks["chunk_words"].to_numpy()[:len(cdf)]
    # Word-weighted document means: a short trailing passage should not count as much as a
    # full-length one.
    def _wmean(g, col):
        w = g["chunk_words"].clip(lower=1)
        v = g[col]
        ok = v.notna()
        return float((v[ok] * w[ok]).sum() / w[ok].sum()) if ok.any() else np.nan
    agg = pd.DataFrame([{
        "doc_id": d,
        "sentiment_score": _wmean(g, "sentiment_score"),
        "p_positive": _wmean(g, "p_positive"),
        "p_neutral": _wmean(g, "p_neutral"),
        "p_negative": _wmean(g, "p_negative"),
        "sentiment_sd": float(g["sentiment_score"].std()) if len(g) > 1 else np.nan,
        "sentiment_chunks": int(len(g)),
    } for d, g in cdf.groupby("doc_id", sort=False)])
    agg["sentiment"] = agg[["p_positive","p_neutral","p_negative"]].idxmax(axis=1).str.replace("p_", "", regex=False)
    return cdf, agg



def entity_documents(docs: pd.DataFrame, max_words: int = 180) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Optional multilingual NER on non-overlapping chunks to avoid overlap double-counting."""
    pipe = get_ner_pipeline()
    rows = []
    for r in docs.itertuples():
        # Reuse chunker via a small temporary document set, but force zero overlap for NER.
        tmp = build_chunks(pd.DataFrame([{"doc_id": str(r.doc_id), "text": str(r.text), "language": str(r.language)}]), max_words, 0)
        for c in tmp.itertuples():
            try:
                out = pipe(str(c.chunk_text), truncation=True, max_length=384)
            except TypeError:
                out = pipe(str(c.chunk_text))
            for item in out or []:
                surface = str(item.get("word", "")).replace("▁", " ").strip()
                etype = str(item.get("entity_group", item.get("entity", ""))).replace("B-", "").replace("I-", "")
                score = float(item.get("score", np.nan))
                if surface and etype and (not np.isfinite(score) or score >= 0.45):
                    rows.append({"doc_id": str(r.doc_id), "chunk_id": str(c.chunk_id), "entity": surface,
                                 "entity_type": etype, "score": score})
    mentions = pd.DataFrame(rows)
    return mentions, aggregate_entities(mentions)


def build_analysis_dataset(R: Dict[str, object]) -> pd.DataFrame:
    df = R["docs"].copy()
    if "frames" in R:
        df = df.merge(R["frames"], on="doc_id", how="left")
    if "doc_stats" in R:
        df = df.merge(R["doc_stats"].drop(columns=["source_file","language"], errors="ignore"), on="doc_id", how="left")
    tm = R.get("topics", {})
    for key in ["lda_doc_topic", "nmf_doc_topic"]:
        tdf = tm.get(key) if isinstance(tm, dict) else None
        if isinstance(tdf, pd.DataFrame) and len(tdf) == len(df):
            df = pd.concat([df.reset_index(drop=True), tdf.reset_index(drop=True)], axis=1)
    # Researcher-defined measures computed in the dictionary and concept tabs become
    # available as outcomes in the grouped, temporal, comparison and regression tabs.
    for extra_key in ["pm_custom_frames", "pm_concepts_docs"]:
        extra = st.session_state.get(extra_key)
        if isinstance(extra, pd.DataFrame) and not extra.empty and "doc_id" in extra.columns:
            extra = extra.copy()
            extra["doc_id"] = extra["doc_id"].astype(str)
            new_cols = [c for c in extra.columns if c == "doc_id" or c not in df.columns]
            df["doc_id"] = df["doc_id"].astype(str)
            df = df.merge(extra[new_cols], on="doc_id", how="left")
    return df


def actor_topic_network_figure(edges: pd.DataFrame):
    import networkx as nx
    if edges is None or edges.empty:
        return None
    G = nx.Graph()
    for r in edges.itertuples():
        a = f"A::{r.entity}"
        t = f"T::{r.topic}"
        G.add_node(a, label=str(r.entity), kind="actor", entity_type=str(r.entity_type))
        G.add_node(t, label=str(r.topic), kind="topic", entity_type="topic")
        G.add_edge(a, t, weight=float(r.weight))
    pos = nx.spring_layout(G, seed=42, weight="weight", k=max(0.35, 1.8 / max(len(G), 2) ** 0.5))
    edge_x, edge_y = [], []
    for u, v, d in G.edges(data=True):
        x0, y0 = pos[u]; x1, y1 = pos[v]
        edge_x += [x0, x1, None]; edge_y += [y0, y1, None]
    edge_trace = go.Scatter(x=edge_x, y=edge_y, mode="lines", hoverinfo="none", line=dict(width=0.8))
    actor_x=[]; actor_y=[]; actor_text=[]; actor_hover=[]; actor_size=[]
    topic_x=[]; topic_y=[]; topic_text=[]; topic_hover=[]; topic_size=[]
    deg = dict(G.degree(weight="weight"))
    for n, d in G.nodes(data=True):
        x, y = pos[n]
        size = 10 + 4*np.sqrt(max(deg.get(n, 0), 0))
        if d["kind"] == "actor":
            actor_x.append(x); actor_y.append(y); actor_text.append(d["label"]); actor_hover.append(f"{d['label']} · {d['entity_type']}"); actor_size.append(size)
        else:
            topic_x.append(x); topic_y.append(y); topic_text.append(d["label"]); topic_hover.append(d["label"]); topic_size.append(size)
    actor_trace = go.Scatter(x=actor_x, y=actor_y, mode="markers+text", text=actor_text, textposition="top center",
                             hovertext=actor_hover, hoverinfo="text", marker=dict(size=actor_size), name="Actors/entities")
    topic_trace = go.Scatter(x=topic_x, y=topic_y, mode="markers+text", text=topic_text, textposition="bottom center",
                             hovertext=topic_hover, hoverinfo="text", marker=dict(size=topic_size, symbol="diamond"), name="Topics")
    fig = go.Figure([edge_trace, actor_trace, topic_trace])
    fig.update_layout(height=700, showlegend=True, xaxis=dict(visible=False), yaxis=dict(visible=False), margin=dict(l=10,r=10,t=40,b=10))
    return fig


def make_html_report(R: Dict[str, object], params: Dict[str, object], lang: str = "en") -> bytes:
    master = R["master_docs"]
    stats = R["doc_stats"]
    meth = methodology_table(params, R)
    kw = R.get("keywords", pd.DataFrame()).head(25)
    frames = R.get("frames", pd.DataFrame())
    frame_cols = [c for c in frames.columns if c.endswith("_per_1000")] if isinstance(frames, pd.DataFrame) else []
    frame_summary = frames[frame_cols].mean().sort_values(ascending=False).rename("mean_per_1000").reset_index().rename(columns={"index":"frame"}) if frame_cols else pd.DataFrame()
    cdocs = st.session_state.get("pm_concepts_docs")
    concept_summary = (cdocs.drop(columns=["doc_id"]).mean().rename("corpus_mean").reset_index().rename(columns={"index": "measure"})
                       if isinstance(cdocs, pd.DataFrame) and not cdocs.empty else pd.DataFrame())
    custom = st.session_state.get("pm_custom_frames")
    if isinstance(custom, pd.DataFrame) and not custom.empty:
        frames = custom
        frame_cols = [c for c in frames.columns if c.endswith("_per_1000")]
        frame_summary = frames[frame_cols].mean().sort_values(ascending=False).rename("mean_per_1000").reset_index().rename(columns={"index":"frame"}) if frame_cols else pd.DataFrame()
    def table(df):
        return df.to_html(index=False, border=0, classes="pm-table", escape=True) if isinstance(df,pd.DataFrame) and not df.empty else "<p>Not available.</p>"
    body = f"""<!doctype html><html><head><meta charset='utf-8'><title>PerspectiveMapper v3.3 report</title>
    <style>body{{font-family:Arial,sans-serif;max-width:1100px;margin:40px auto;line-height:1.45;color:#222}}h1,h2{{margin-top:1.5em}}.metrics{{display:flex;gap:30px;flex-wrap:wrap}}.metric{{padding:12px 18px;border:1px solid #ddd;border-radius:8px}}table{{border-collapse:collapse;width:100%;font-size:14px}}th,td{{border-bottom:1px solid #ddd;padding:7px;text-align:left}}.note{{background:#f5f5f5;padding:12px;border-radius:6px}}</style>
    </head><body><h1>PerspectiveMapper v3.3 — Research Report</h1>
    <p class='note'>Generated from the analysed corpus. NLP outputs are measurements with model uncertainty and require substantive validation.</p>
    <div class='metrics'><div class='metric'><b>Documents</b><br>{len(master)}</div><div class='metric'><b>Words</b><br>{int(stats['words'].sum())}</div><div class='metric'><b>Languages</b><br>{master['language'].nunique()}</div></div>
    <h2>Methodology</h2>{table(meth)}<h2>Top TF-IDF terms</h2>{table(kw)}
    <h2>Topics (NMF)</h2>{table(R.get("topics", {}).get("nmf_topics"))}<h2>Topics (LDA)</h2>{table(R.get("topics", {}).get("lda_topics"))}
    <h2>Mean dictionary indicators</h2>{table(frame_summary)}
    <h2>Concept projection (corpus means)</h2>{table(concept_summary)}
    <h2>Concept anchors</h2>{table(st.session_state.get("pm_concepts_info"))}
    <h2>Reproducibility note</h2><p>Document-level inference is appropriate only when documents are defensible independent analytical units. For nested or repeated data, use a corresponding multilevel/panel design.</p>
    </body></html>"""
    return body.encode("utf-8")


def publication_package(R: Dict[str, object], params: Dict[str, object]) -> bytes:
    """ZIP with 300-dpi PNG/SVG figures, CSV tables, HTML report and methodology."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", compression=zipfile.ZIP_DEFLATED) as z:
        master = R["master_docs"]
        meth = methodology_table(params, R)
        z.writestr("tables/documents.csv", master.drop(columns=["text"], errors="ignore").to_csv(index=False))
        z.writestr("tables/document_statistics.csv", R["doc_stats"].to_csv(index=False))
        z.writestr("tables/keywords_tfidf.csv", R.get("keywords", pd.DataFrame()).to_csv(index=False))
        z.writestr("tables/frames.csv", R.get("frames", pd.DataFrame()).to_csv(index=False))
        z.writestr("methodology.csv", meth.to_csv(index=False))
        z.writestr("report.html", make_html_report(R, params))
        if "entities" in R: z.writestr("tables/entities.csv", R["entities"].to_csv(index=False))
        for _key, _fname in [("pm_custom_frames", "dictionary_indicators"), ("pm_custom_frames_audit", "dictionary_audit"),
                             ("pm_concepts_docs", "concepts_documents"), ("pm_concepts_info", "concept_anchors"),
                             ("pm_topic_k_diag", "topic_k_diagnostics"), ("pm_grouped_summary", "grouped_summary")]:
            _df = st.session_state.get(_key)
            if isinstance(_df, pd.DataFrame) and not _df.empty:
                z.writestr(f"tables/{_fname}.csv", _df.to_csv(index=False))
        if "sentiment_docs" in R: z.writestr("tables/sentiment_documents.csv", R["sentiment_docs"].to_csv(index=False))
        tm = R.get("topics", {})
        for key in ["lda_topics","lda_doc_topic","nmf_topics","nmf_doc_topic"]:
            if isinstance(tm.get(key), pd.DataFrame): z.writestr(f"tables/{key}.csv", tm[key].to_csv(index=False))
        figures=[]
        kw = R.get("keywords", pd.DataFrame()).head(20)
        if isinstance(kw,pd.DataFrame) and not kw.empty:
            fig, ax = plt.subplots(figsize=(8,6)); q=kw.sort_values("mean_tfidf"); ax.barh(q["term"], q["mean_tfidf"]); ax.set_xlabel("Mean TF-IDF"); ax.set_title("Top distinctive terms"); fig.tight_layout(); figures.append(("figure_01_tfidf",fig))
        if "sentiment_docs" in R and not R["sentiment_docs"].empty:
            sd=R["sentiment_docs"].sort_values("sentiment_score")
            fig, ax = plt.subplots(figsize=(9,5)); ax.bar(sd["doc_id"].astype(str), sd["sentiment_score"]); ax.axhline(0,linewidth=.8); ax.set_ylabel("P(positive) − P(negative)"); ax.set_title("Document-level sentiment"); ax.tick_params(axis='x',rotation=90); fig.tight_layout(); figures.append(("figure_02_sentiment",fig))
        frames=st.session_state.get("pm_custom_frames") if isinstance(st.session_state.get("pm_custom_frames"), pd.DataFrame) else R.get("frames",pd.DataFrame()); fcols=[c for c in frames.columns if c.endswith("_per_1000")] if isinstance(frames,pd.DataFrame) else []
        if fcols:
            means=frames[fcols].mean().sort_values()
            fig, ax = plt.subplots(figsize=(8,5)); ax.barh([x.replace("frame_","").replace("_per_1000","") for x in means.index],means.values); ax.set_xlabel("Mean hits per 1,000 tokens"); ax.set_title("Framing indicators"); fig.tight_layout(); figures.append(("figure_03_frames",fig))
        tdf=tm.get("nmf_doc_topic") if isinstance(tm,dict) else None
        if isinstance(tdf,pd.DataFrame) and not tdf.empty:
            means=tdf.mean().sort_values()
            fig, ax=plt.subplots(figsize=(8,5)); ax.barh(means.index,means.values); ax.set_xlabel("Mean topic weight"); ax.set_title("NMF topic prevalence"); fig.tight_layout(); figures.append(("figure_04_topics",fig))
        cdocs = st.session_state.get("pm_concepts_docs")
        if isinstance(cdocs, pd.DataFrame) and not cdocs.empty:
            ccols = [c for c in cdocs.columns if c.endswith("_share")]
            if ccols:
                means = cdocs[ccols].mean().sort_values()
                fig, ax = plt.subplots(figsize=(8,5)); ax.barh([c[len("concept_"):-len("_share")] for c in means.index], means.values); ax.set_xlabel("Mean share of words assigned to concept"); ax.set_title("Concept projection"); fig.tight_layout(); figures.append(("figure_05_concepts",fig))
        for name,fig in figures:
            png=io.BytesIO(); svg=io.BytesIO(); fig.savefig(png,format="png",dpi=300,bbox_inches="tight"); fig.savefig(svg,format="svg",bbox_inches="tight"); plt.close(fig)
            z.writestr(f"figures/{name}.png",png.getvalue()); z.writestr(f"figures/{name}.svg",svg.getvalue())
    return bio.getvalue()

def read_unstructured(name: str, data: bytes, clean: bool = True, strip_notes: bool = False) -> str:
    lower = name.lower()
    if lower.endswith(".txt"):
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        return clean_extracted_text(text, strip_notes=strip_notes) if clean else text
    if lower.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        blocks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if vals: blocks.append(" | ".join(vals))
        text = "\n\n".join(blocks)
        return clean_extracted_text(text, strip_notes=strip_notes) if clean else text
    if lower.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        if not clean:
            return "\n".join(pages)
        repeated = find_repeated_page_lines(pages)
        return clean_extracted_text("\n".join(pages), repeated, strip_notes=strip_notes)
    return ""


@st.cache_data(show_spinner=False, max_entries=64)
def read_table_bytes(name: str, data: bytes, sheet: Optional[str] = None) -> pd.DataFrame:
    bio = io.BytesIO(data)
    if name.lower().endswith(".csv"):
        # Let pandas sniff the delimiter; try UTF-8 (with BOM) first, then Windows/Latin-1.
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                bio.seek(0)
                return pd.read_csv(bio, sep=None, engine="python", encoding=enc)
            except UnicodeDecodeError:
                continue
    return pd.read_excel(bio, sheet_name=sheet or 0)


def read_table(upload, sheet: Optional[str] = None) -> pd.DataFrame:
    return read_table_bytes(upload.name, upload.getvalue(), sheet)


@st.cache_data(show_spinner=False, max_entries=64)
def excel_sheet_names(name: str, data: bytes) -> List[str]:
    try:
        return pd.ExcelFile(io.BytesIO(data)).sheet_names
    except Exception:
        return []


def excel_sheets(upload) -> List[str]:
    if not upload.name.lower().endswith((".xlsx", ".xls")):
        return []
    return excel_sheet_names(upload.name, upload.getvalue())


def _natural_text_order(columns: List[str]) -> List[str]:
    """Sort a split-text family as texto, texto_2, texto_3, ... rather than lexicographically."""
    def key(c: str):
        m = re.match(r"^(.*?)(?:_(\d+))?$", str(c))
        stem = m.group(1) if m else str(c)
        suffix = int(m.group(2)) if m and m.group(2) else 1
        return (stem.lower(), suffix)
    return sorted([str(c) for c in columns], key=key)


def candidate_text_columns(df: pd.DataFrame) -> List[str]:
    """Detect one text column or a split family such as text, text_2, ... .

    Excel cells are limited to 32,767 characters, so long documents are often
    distributed across sequential columns. This detector scores both single
    columns and such families by the amount of textual content they contain.
    """
    candidates = [str(c) for c in df.columns if df[c].dtype == "object" or pd.api.types.is_string_dtype(df[c])]
    if not candidates:
        return [str(df.columns[0])]

    # Single-column score: mean non-empty characters per row.
    single_scores = {}
    for c in candidates:
        vals = df[c].fillna("").astype(str).head(500)
        single_scores[c] = float(vals.str.len().mean()) if len(vals) else 0.0
    best_single = max(single_scores, key=single_scores.get)
    best_cols = [best_single]
    best_score = single_scores[best_single]

    # Candidate split families: <stem>, <stem>_2, <stem>_3, ...
    families: Dict[str, List[str]] = {}
    for c in candidates:
        m = re.match(r"^(.*?)(?:_(\d+))?$", c)
        if not m:
            continue
        stem = m.group(1)
        families.setdefault(stem, []).append(c)

    for stem, cols in families.items():
        cols = _natural_text_order(cols)
        # Require at least two columns and an actual numbered continuation.
        if len(cols) < 2 or not any(re.match(rf"^{re.escape(stem)}_\d+$", c) for c in cols):
            continue
        # Avoid treating short metadata series as split documents.
        total_chars = pd.Series(0.0, index=df.index[:500])
        for c in cols:
            total_chars = total_chars.add(df.loc[total_chars.index, c].fillna("").astype(str).str.len(), fill_value=0)
        score = float(total_chars.mean()) if len(total_chars) else 0.0
        if score > best_score * 1.05 and score >= 200:
            best_cols, best_score = cols, score

    return best_cols


def candidate_text_column(df: pd.DataFrame) -> str:
    """Backward-compatible first candidate."""
    return candidate_text_columns(df)[0]


@st.cache_data(show_spinner="Reading corpus…", max_entries=8)
def _prepare_corpus_cached(files: Tuple[Tuple[str, bytes], ...], configs_json: str,
                           clean_text: bool, strip_notes: bool) -> pd.DataFrame:
    """Read, clean and language-tag the corpus once per (files, configuration).

    Cached so that ordinary widget interactions do not re-parse PDFs or re-run language
    detection on every Streamlit rerun.
    """
    configs = json.loads(configs_json)
    rows = []
    used_ids = set()
    for up_name, up_data in files:
        lower = up_name.lower()
        cfg = configs.get(up_name, {})
        if lower.endswith((".csv", ".xlsx", ".xls")):
            df = read_table_bytes(up_name, up_data, cfg.get("sheet"))
            text_cols = cfg.get("text_cols") or ([cfg.get("text_col")] if cfg.get("text_col") else [])
            text_cols = [c for c in text_cols if c in df.columns]
            if not text_cols:
                continue
            id_col = cfg.get("id_col")
            meta_cols = cfg.get("meta_cols", [])
            for ridx, row in df.iterrows():
                parts = []
                for c in text_cols:
                    if c in row.index and not pd.isna(row[c]):
                        value = str(row[c]).strip()
                        if value:
                            parts.append(value)
                text = "\n\n".join(parts).strip()
                if clean_text:
                    text = clean_extracted_text(text, strip_notes=strip_notes)
                if not text:
                    continue
                base_id = str(row[id_col]) if id_col and not pd.isna(row[id_col]) else f"{Path(up_name).stem}_{ridx+1}"
                doc_id = base_id
                n = 2
                while doc_id in used_ids:
                    doc_id = f"{base_id}_{n}"; n += 1
                used_ids.add(doc_id)
                rec = {"doc_id": doc_id, "source_file": up_name, "text": text}
                for c in meta_cols:
                    rec[str(c)] = row[c] if c in row.index else np.nan
                rows.append(rec)
        else:
            text = read_unstructured(up_name, up_data, clean_text, strip_notes).strip()
            if not text:
                continue
            base_id = Path(up_name).stem
            doc_id = base_id; n = 2
            while doc_id in used_ids:
                doc_id = f"{base_id}_{n}"; n += 1
            used_ids.add(doc_id)
            rows.append({"doc_id": doc_id, "source_file": up_name, "text": text})
    docs = pd.DataFrame(rows)
    if docs.empty:
        return docs
    docs["language"] = docs["text"].map(safe_language)
    return docs


def prepare_corpus(uploads, configs, clean_text: bool = True, strip_notes: bool = False) -> pd.DataFrame:
    files = tuple((u.name, u.getvalue()) for u in uploads)
    return _prepare_corpus_cached(files, json.dumps(configs, default=str, sort_keys=True), bool(clean_text), bool(strip_notes)).copy()


@st.cache_data(show_spinner=False, max_entries=8)
def cached_topic_models(units: Tuple[str, ...], n_topics: int, max_features: int, random_state: int) -> Dict[str, object]:
    return topic_models(list(units), n_topics, max_features, random_state)


def resolve_topic_unit(topic_unit: str, docs: pd.DataFrame) -> str:
    """'auto' fits topics on passages when documents are few or long (e.g. encyclicals)."""
    if topic_unit in {"document", "chunk"}:
        return topic_unit
    words = docs["text"].astype(str).str.split().str.len()
    return "chunk" if (len(docs) < 30 or float(words.median()) > 1500) else "document"


def run_pipeline(docs: pd.DataFrame, chunk_words: int, overlap_words: int, n_topics: int,
                 max_features: int, do_semantic: bool, do_sentiment: bool, do_entities: bool,
                 auto_k: bool, k_manual: int, random_state: int, lemmatize: bool = False,
                 extra_stopwords: Optional[List[str]] = None, topic_unit: str = "auto",
                 progress=None) -> Dict[str, object]:
    def step(frac: float, msg: str):
        if progress is not None:
            progress.progress(min(1.0, frac), text=msg)

    def sub(lo: float, hi: float, label: str):
        return lambda f, extra="": step(lo + (hi - lo) * f, f"{label} {extra}")

    res: Dict[str, object] = {"docs": docs.copy()}
    step(0.02, "Chunking…")
    chunks = build_chunks(docs, chunk_words, overlap_words)
    res["chunks"] = chunks
    stats_df = document_statistics(docs)
    res["doc_stats"] = stats_df
    step(0.08, "Lexical preprocessing…")
    extra = list(extra_stopwords or [])
    cleaned = [clean_for_lexical(r.text, r.language, extra, lemmatize) for r in docs.itertuples()]
    res["cleaned_docs"] = cleaned
    res["keywords"] = top_tfidf_terms(cleaned, 40, max_features)
    res["collocations"] = collocations_pmi(cleaned, min_count=max(2, int(np.ceil(len(docs)/8))), top_n=50)

    step(0.15, "Topic models…")
    unit = resolve_topic_unit(topic_unit, docs)
    res["topic_unit"] = unit
    if unit == "chunk":
        # Non-overlapping passages, so no text is counted twice in the topic model.
        tchunks = build_chunks(docs, chunk_words, 0)
        lang_map = dict(zip(docs["doc_id"].astype(str), docs["language"]))
        tunits = [clean_for_lexical(r.chunk_text, lang_map.get(str(r.doc_id), "und"), extra, lemmatize) for r in tchunks.itertuples()]
        step(0.17, f"Topic models… LDA + NMF on {len(tunits):,} passages")
        tm = dict(cached_topic_models(tuple(tunits), n_topics, max_features, random_state))
        doc_ids = docs["doc_id"].astype(str).tolist()
        for m in ["lda", "nmf"]:
            if isinstance(tm.get(f"{m}_doc_topic"), pd.DataFrame):
                tm[f"{m}_unit_topic"] = pd.concat([tchunks[["chunk_id", "doc_id"]].reset_index(drop=True), tm[f"{m}_doc_topic"]], axis=1)
                tm[f"{m}_doc_topic"] = aggregate_unit_topics(tm[f"{m}_doc_topic"], tchunks["doc_id"], doc_ids, tchunks["chunk_words"])
        res["topic_units"] = tunits
    else:
        tm = dict(cached_topic_models(tuple(cleaned), n_topics, max_features, random_state))
        res["topic_units"] = cleaned
    for m in ["lda", "nmf"]:
        if tm.get(f"{m}_top_terms"):
            coh = topic_coherence_npmi(tm[f"{m}_top_terms"], res["topic_units"], 10)
            tm[f"{m}_topics"] = tm[f"{m}_topics"].merge(coh, on="topic", how="left")
            tm[f"{m}_diversity"] = topic_diversity(tm[f"{m}_top_terms"], 25)
    res["topics"] = tm
    res["frames"] = frame_indicators(docs, DEFAULT_FRAMES)

    if do_semantic and len(chunks):
        step(0.35, "Multilingual embeddings…")
        try:
            chunk_emb = embed_texts(tuple(chunks["chunk_text"].astype(str).tolist()), sub(0.35, 0.62, "Multilingual embeddings…"))
            doc_emb = aggregate_embeddings(chunk_emb, chunks, docs["doc_id"].astype(str).tolist())
            sem_units = chunk_emb if len(chunks) >= 4 else doc_emb
            sem = semantic_clustering(sem_units, auto_k, k_manual, random_state)
            res["chunk_embeddings"] = chunk_emb
            res["doc_embeddings"] = doc_emb
            res["semantic"] = sem
            res["semantic_unit"] = "chunk" if len(chunks) >= 4 else "document"
            # doc_emb rows are unit-length (see aggregate_embeddings), so this is cosine similarity.
            res["doc_similarity"] = pd.DataFrame(
                np.clip(doc_emb @ doc_emb.T, -1, 1), index=docs["doc_id"], columns=docs["doc_id"]
            )
            res["hierarchy"] = hierarchy_from_embeddings(doc_emb)
        except Exception as e:
            res["semantic_error"] = str(e)

    if do_sentiment and len(chunks):
        step(0.65, "Sentiment…")
        try:
            sc, sd = sentiment_chunks(chunks, sub(0.65, 0.84, "Sentiment…"))
            res["sentiment_chunks"] = sc
            res["sentiment_docs"] = sd
            res["docs"] = res["docs"].merge(sd, on="doc_id", how="left")
        except Exception as e:
            res["sentiment_error"] = str(e)

    if do_entities:
        step(0.85, "Named entities…")
        try:
            mentions, entities = entity_documents(docs)
            res["entity_mentions"] = mentions
            res["entities"] = entities
        except Exception as e:
            res["entities_error"] = str(e)

    # Add statistics and frames to master document table.
    master = res["docs"].merge(stats_df.drop(columns=["source_file","language"], errors="ignore"), on="doc_id", how="left")
    master = master.merge(res["frames"], on="doc_id", how="left")
    res["master_docs"] = master
    step(1.0, "Done")
    return res


def methodology_table(params: Dict[str, object], results: Dict[str, object]) -> pd.DataFrame:
    rows = [
        ("PerspectiveMapper version", "3.3 Doctrine & Concepts Edition"),
        ("Python", sys.version.split()[0]),
        ("Embedding model", EMBEDDING_MODEL if params.get("semantic") else "disabled"),
        ("Sentiment model", SENTIMENT_MODEL if params.get("sentiment") else "disabled"),
        ("Named-entity model", NER_MODEL if params.get("entities") else "disabled"),
        ("Text cleaning", ("PDF/DOCX repair: page numbers, running headers/footers, line-end hyphenation, hard wraps"
                           + ("; final notes section and [n] markers removed" if params.get("strip_notes") else ""))
                          if params.get("clean_text") else "none (raw extraction)"),
        ("Lemmatisation", "simplemma dictionary lemmatiser (lexical and topic models only)" if params.get("lemmatize") else "none (surface forms)"),
        ("Additional stopwords", ", ".join(params.get("extra_stopwords", [])) or "none"),
        ("Topic-model unit", {"chunk": "non-overlapping passages; document mixtures = word-weighted passage means",
                              "document": "whole documents"}.get(results.get("topic_unit"), str(results.get("topic_unit")))),
        ("LDA estimation", "online variational Bayes (batch size 512, 10 passes) above 1,000 units; batch variational Bayes (30 iterations) otherwise"),
        ("Topic quality", "NPMI coherence of top-10 terms (unit co-occurrence); topic diversity of top-25 terms"),
        ("Chunk size words", params.get("chunk_words")),
        ("Chunk overlap words", params.get("overlap_words")),
        ("Requested topics", params.get("n_topics")),
        ("Maximum lexical features", params.get("max_features")),
        ("Random seed", params.get("random_state")),
        ("Keyword filter", ", ".join(params.get("keyword_terms", [])) if params.get("keyword_terms") else "none"),
        ("Keyword filter mode", params.get("keyword_mode", "ANY") if params.get("keyword_terms") else "not applicable"),
        ("Keyword filter scope", params.get("keyword_scope", "text") if params.get("keyword_terms") else "not applicable"),
        ("Semantic cluster selection", "silhouette" if params.get("auto_k") else f"manual k={params.get('k_manual')}"),
        ("Semantic clustering space", "full normalized embedding space; PCA only for 2D visualisation"),
        ("Hierarchical clustering", "average linkage on cosine distances"),
        ("Sentiment aggregation", "chunk probabilities aggregated to document level (word-weighted mean)"),
        ("Concept projection", "cosine similarity of passages to researcher-defined anchor centroids; word-weighted document means, corpus-centred scores and assignment shares"),
        ("Framing indicators", "dictionary hits per 1,000 tokens; descriptive, not automated bias classification"),
        ("Supervised stance/framing", "multilingual document embeddings + class-balanced logistic regression; stratified cross-validation on user-coded labels"),
        ("Bootstrap", "percentile bootstrap resampling at document level"),
    ]
    return pd.DataFrame(rows, columns=["parameter", "value"])


# ---------- Sidebar: language before gate ----------
with st.sidebar:
    chosen_ui = st.selectbox("Language / Idioma / Lingua", list(UI_LANGUAGES.keys()), index=0)
lang = UI_LANGUAGES[chosen_ui]
optional_gate(lang)

# ---------- Header ----------
logo = ROOT / "assets" / "logo.png"
c1, c2 = st.columns([1, 7])
with c1:
    if logo.exists(): st.image(str(logo), width=96)
with c2:
    st.title("PerspectiveMapper v3.3")
    st.caption(tr(lang, "subtitle"))

with st.sidebar:
    st.header(f"⚙️ {tr(lang, 'settings')}")
    uploads = st.file_uploader(
        tr(lang, "upload"), type=["txt","docx","pdf","csv","xlsx","xls"],
        accept_multiple_files=True, help=tr(lang, "upload_help")
    )

configs: Dict[str, dict] = {}
if uploads:
    tabular = [u for u in uploads if u.name.lower().endswith((".csv", ".xlsx", ".xls"))]
    if tabular:
        with st.sidebar.expander(tr(lang, "structured"), expanded=True):
            for up in tabular:
                st.markdown(f"**{up.name}**")
                sheets = excel_sheets(up)
                sheet = st.selectbox(tr(lang, "sheet"), sheets, key=f"sheet_{up.name}") if sheets else None
                try:
                    preview = read_table(up, sheet).head(50)
                    cols = [str(c) for c in preview.columns]
                    auto_cols = candidate_text_columns(preview)
                    text_cols = st.multiselect(
                        tr(lang, "text_columns"), cols, default=[c for c in auto_cols if c in cols], key=f"text_{up.name}",
                        help=tr(lang, "text_columns_help")
                    )
                    if len(text_cols) > 1:
                        st.caption(tr(lang, "text_columns_concat") + ": " + " → ".join(text_cols))
                    if text_cols:
                        _char_total = int(sum(preview[c].fillna("").astype(str).str.len().sum() for c in text_cols if c in preview.columns))
                        _nonempty_rows = int(preview[text_cols].fillna("").astype(str).apply(lambda row: any(v.strip() for v in row), axis=1).sum())
                        st.caption(f"{tr(lang,'text_preview')}: {_nonempty_rows} · {tr(lang,'characters_preview')}: {_char_total:,}")
                    id_options = ["—"] + cols
                    preferred_ids = ["doc_id", "document_id", "id", "nombre_enciclica", "title", "titulo", "archivo_fuente"]
                    auto_id = next((c for c in preferred_ids if c in cols and c not in text_cols), "—")
                    id_col_raw = st.selectbox(tr(lang, "id_column"), id_options, index=id_options.index(auto_id), key=f"id_{up.name}")
                    meta_default = [c for c in cols if c not in text_cols and c != id_col_raw][:8]
                    meta_cols = st.multiselect(tr(lang, "metadata"), [c for c in cols if c not in text_cols], default=meta_default, key=f"meta_{up.name}")
                    configs[up.name] = {"sheet": sheet, "text_cols": text_cols, "id_col": None if id_col_raw == "—" else id_col_raw, "meta_cols": meta_cols}
                except Exception as e:
                    st.warning(f"{up.name}: {e}")

with st.sidebar.expander(f"🧹 {tr(lang, 'preprocessing')}", expanded=False):
    clean_text = st.checkbox(tr(lang, "clean_text"), value=True, help=tr(lang, "clean_text_help"))
    strip_notes = st.checkbox(tr(lang, "strip_notes"), value=False, help=tr(lang, "strip_notes_help"), disabled=not clean_text)
    lemmatize = st.checkbox(tr(lang, "lemmatize"), value=lemmatizer_available(), disabled=not lemmatizer_available(),
                            help=tr(lang, "lemmatize_help"))
    extra_stop_raw = st.text_area(tr(lang, "extra_stopwords"), value="", height=70,
                                  placeholder=tr(lang, "extra_stopwords_placeholder"), help=tr(lang, "extra_stopwords_help"))
    extra_stopwords = sorted({w.strip().lower() for w in re.split(r"[,;\n]+", extra_stop_raw) if w.strip()})

with st.sidebar:
    st.subheader(tr(lang, "analysis_settings"))
    chunk_words = st.slider(tr(lang, "chunk_words"), 60, 300, 140, 10)
    overlap_words = st.slider(tr(lang, "overlap_words"), 0, 60, 20, 5)
    n_topics = st.slider(tr(lang, "topics"), 2, 30, 8)
    topic_unit = st.selectbox(tr(lang, "topic_unit"), ["auto", "chunk", "document"],
                              format_func=lambda x: tr(lang, f"topic_unit_{x}"), help=tr(lang, "topic_unit_help"))
    max_features = st.slider(tr(lang, "max_features"), 500, 12000, 5000, 500)
    do_semantic = st.checkbox(tr(lang, "semantic"), value=True)
    do_sentiment = st.checkbox(tr(lang, "sentiment"), value=False, help=tr(lang, "sentiment_help"))
    do_entities = st.checkbox(tr(lang, "entities_option"), value=False, help=tr(lang, "entities_help"))
    auto_k = st.checkbox(tr(lang, "auto_k"), value=True)
    k_manual = st.slider(tr(lang, "clusters"), 2, 12, 4, disabled=auto_k)
    random_state = st.number_input("Random seed", value=42, step=1)
    run_clicked = False

if not uploads:
    st.info(tr(lang, "need_files"))
    st.stop()

try:
    docs_now = prepare_corpus(uploads, configs, clean_text, strip_notes)
except Exception as e:
    st.error(f"{tr(lang, 'error')}: {e}")
    st.stop()

if docs_now.empty:
    st.error("No readable text was found in the uploaded files.")
    st.stop()

# ---------- Global keyword filtering (document remains the unit of analysis) ----------
docs_before_filter = docs_now.copy()
metadata_for_filter = [c for c in docs_now.columns if c not in {"doc_id", "text", "language"}]
with st.sidebar.expander(f"🔎 {tr(lang, 'keyword_filter')}", expanded=False):
    keyword_raw = st.text_area(
        tr(lang, "keyword_terms"), value="", height=90,
        placeholder=tr(lang, "keyword_placeholder"),
        help=tr(lang, "keyword_help"), key="pm_keyword_raw"
    )
    keyword_terms = parse_keyword_terms(keyword_raw)
    cfa, cfb = st.columns(2)
    with cfa:
        keyword_mode = st.selectbox(
            tr(lang, "keyword_mode"), ["ANY", "ALL"],
            format_func=lambda x: tr(lang, "keyword_any") if x == "ANY" else tr(lang, "keyword_all"),
            key="pm_keyword_mode"
        )
    with cfb:
        keyword_scope = st.selectbox(
            tr(lang, "keyword_scope"), ["text", "both", "metadata"],
            format_func=lambda x: {"text": tr(lang,"keyword_scope_text"), "both": tr(lang,"keyword_scope_both"), "metadata": tr(lang,"keyword_scope_metadata")}[x],
            key="pm_keyword_scope"
        )
    accent_insensitive = st.checkbox(tr(lang, "accent_insensitive"), value=True, key="pm_accent_insensitive")

if keyword_terms:
    docs_now, keyword_audit = filter_documents_by_keywords(
        docs_now, keyword_terms, mode=keyword_mode, scope=keyword_scope,
        metadata_cols=metadata_for_filter, accent_insensitive=accent_insensitive,
    )
else:
    docs_now, keyword_audit = filter_documents_by_keywords(docs_now, [])

if docs_now.empty:
    st.error(tr(lang, "keyword_no_matches"))
    st.stop()

# ---------- Global result grouping ----------
_group_excluded = {"doc_id", "text", "language", "keyword_matches", "keyword_match_count"}
group_candidates = []
for c in docs_now.columns:
    if c in _group_excluded:
        continue
    nun = docs_now[c].nunique(dropna=True)
    if 2 <= nun <= min(100, max(30, len(docs_now))):
        group_candidates.append(c)
preferred_group_names = ["papa", "pope", "pontiff", "group", "grupo", "country", "pais", "país", "source", "newspaper", "periodico", "periódico"]
auto_group = next((c for c in preferred_group_names if c in group_candidates), None)
with st.sidebar.expander(f"🧩 {tr(lang, 'result_grouping')}", expanded=bool(auto_group)):
    group_options = ["—"] + group_candidates
    if st.session_state.get("pm_result_group") not in group_options:
        st.session_state.pop("pm_result_group", None)
    group_index = group_options.index(auto_group) if auto_group in group_options else 0
    result_group_raw = st.selectbox(tr(lang, "group_results_by"), group_options, index=group_index, key="pm_result_group")
    selected_result_group = None if result_group_raw == "—" else result_group_raw
    if selected_result_group:
        counts_preview = docs_now[selected_result_group].value_counts(dropna=False).rename_axis(selected_result_group).reset_index(name="documents")
        st.dataframe(counts_preview, width="stretch", hide_index=True, height=min(250, 38 * (len(counts_preview) + 1)))

with st.sidebar:
    run_clicked = st.button(f"▶️ {tr(lang, 'run')}", type="primary", width="stretch")

st.caption(f"{tr(lang, 'loaded')}: **{len(docs_now)} {tr(lang, 'documents')}**")
if keyword_terms:
    st.caption(f"🔎 {tr(lang, 'keyword_filter_active')}: **{len(docs_now)}/{len(docs_before_filter)}** · " + ", ".join(keyword_terms))
if selected_result_group:
    st.caption(f"🧩 {tr(lang, 'group_results_by')}: **{selected_result_group}**")

params = {
    "chunk_words": int(chunk_words), "overlap_words": int(overlap_words),
    "n_topics": int(n_topics), "max_features": int(max_features),
    "semantic": bool(do_semantic), "sentiment": bool(do_sentiment), "entities": bool(do_entities),
    "auto_k": bool(auto_k), "k_manual": int(k_manual), "random_state": int(random_state),
    "keyword_terms": keyword_terms, "keyword_mode": keyword_mode if keyword_terms else "ANY",
    "keyword_scope": keyword_scope if keyword_terms else "text", "accent_insensitive": bool(accent_insensitive) if keyword_terms else True,
    "clean_text": bool(clean_text), "strip_notes": bool(strip_notes and clean_text), "lemmatize": bool(lemmatize),
    "extra_stopwords": extra_stopwords, "topic_unit": topic_unit,
}
corpus_signature = (tuple((u.name, len(u.getvalue())) for u in uploads), tuple(sorted((k, str(v)) for k,v in params.items())), json.dumps(configs, default=str, sort_keys=True))

if run_clicked:
    for _k in ["pm_group_bootstrap","pm_grouped_summary","pm_custom_frames","pm_custom_frames_audit","pm_temporal_summary","pm_temporal_trend","pm_network_edges","pm_supervised","pm_supervised_info","pm_model_coefficients","pm_model_info","pm_publication_package","pm_excel","pm_topic_k_diag","pm_concepts","pm_concepts_docs","pm_concepts_info"]:
        st.session_state.pop(_k, None)
    _progress = st.progress(0.0, text="Analysing corpus…")
    try:
        st.session_state.pm_results = run_pipeline(docs_now, **{
            "chunk_words": int(chunk_words), "overlap_words": int(overlap_words),
            "n_topics": int(n_topics), "max_features": int(max_features),
            "do_semantic": bool(do_semantic), "do_sentiment": bool(do_sentiment), "do_entities": bool(do_entities),
            "auto_k": bool(auto_k), "k_manual": int(k_manual), "random_state": int(random_state),
            "lemmatize": bool(lemmatize), "extra_stopwords": extra_stopwords, "topic_unit": topic_unit,
            "progress": _progress,
        })
    finally:
        _progress.empty()
    st.session_state.pm_results["keyword_filter_audit"] = keyword_audit
    st.session_state.pm_results["keyword_filter_terms"] = keyword_terms
    st.session_state.pm_signature = corpus_signature
    st.session_state.pm_params = params
    st.success(tr(lang, "analysis_complete"))

if "pm_results" not in st.session_state or st.session_state.get("pm_signature") != corpus_signature:
    st.info(f"▶️ {tr(lang, 'run')}")
    st.stop()

R = st.session_state.pm_results
master = R["master_docs"]
chunks = R["chunks"]

# ---------- Tabs ----------
tab_over, tab_lex, tab_topics, tab_sem, tab_sent, tab_frames, tab_concepts, tab_entities, tab_grouped, tab_groups, tab_temporal, tab_supervised, tab_models, tab_publication, tab_export = st.tabs([
    f"📊 {tr(lang,'overview')}", f"🔤 {tr(lang,'lexical')}", f"🧵 {tr(lang,'topics_tab')}",
    f"🧭 {tr(lang,'semantic_tab')}", f"💬 {tr(lang,'sentiment_tab')}", f"🧷 {tr(lang,'frames')}", f"🎓 {tr(lang,'concepts_tab')}",
    f"👥 {tr(lang,'entities_tab')}", f"🧩 {tr(lang,'grouped_tab')}", f"⚖️ {tr(lang,'groups')}", f"📈 {tr(lang,'temporal_tab')}",
    f"🎯 {tr(lang,'supervised_tab')}", f"📐 {tr(lang,'models')}", f"📝 {tr(lang,'publication_tab')}", f"⬇️ {tr(lang,'export')}"
])

with tab_over:
    st.subheader(tr(lang, "corpus_metrics"))
    stats_df = R["doc_stats"]
    m1,m2,m3,m4 = st.columns(4)
    m1.metric(tr(lang,"n_docs"), f"{len(master):,}")
    m2.metric(tr(lang,"n_words"), f"{int(stats_df['words'].sum()):,}")
    m3.metric(tr(lang,"n_langs"), f"{master['language'].nunique()}")
    m4.metric(tr(lang,"n_chunks"), f"{len(chunks):,}")
    c1,c2 = st.columns([1,2])
    with c1:
        lang_counts = master["language"].value_counts().rename_axis("language").reset_index(name="documents")
        st.plotly_chart(px.bar(lang_counts, x="language", y="documents", title=tr(lang,"language_distribution")), width="stretch")
    with c2:
        st.subheader(tr(lang,"doc_stats"))
        st.dataframe(stats_df, width="stretch", hide_index=True)
    with st.expander(tr(lang,"docs_table")):
        display_cols = [c for c in master.columns if c != "text"]
        st.dataframe(master[display_cols], width="stretch", hide_index=True)

with tab_lex:
    st.markdown(f"<div class='pm-note'>{tr(lang,'lexical_note')}</div>", unsafe_allow_html=True)
    c1,c2 = st.columns(2)
    with c1:
        st.subheader(tr(lang,"top_terms"))
        kw = R["keywords"]
        if not kw.empty:
            st.plotly_chart(px.bar(kw.head(25).sort_values("mean_tfidf"), x="mean_tfidf", y="term", orientation="h"), width="stretch")
            st.dataframe(kw, width="stretch", hide_index=True)
        else: st.info(tr(lang,"not_enough"))
    with c2:
        st.subheader(tr(lang,"collocations"))
        coll = R["collocations"]
        if coll is not None and not coll.empty:
            st.dataframe(coll, width="stretch", hide_index=True)
        else: st.info(tr(lang,"not_enough"))

def render_topic_model(m: str, tm: Dict[str, object]) -> None:
    label = m.upper()
    st.subheader(tr(lang, m))
    if f"{m}_topics" not in tm:
        st.info(tm.get(f"{m}_error", tr(lang, "not_enough")))
        return
    topics_df = tm[f"{m}_topics"]
    k1, k2, k3 = st.columns(3)
    if m == "lda":
        k1.metric(tr(lang, "perplexity"), f"{tm.get('lda_perplexity', np.nan):.1f}")
    else:
        k1.metric(tr(lang, "reconstruction_error"), f"{tm.get('nmf_reconstruction_error', np.nan):.3f}")
    if "npmi" in topics_df.columns:
        k2.metric(tr(lang, "mean_npmi"), f"{topics_df['npmi'].mean():.3f}")
    k3.metric(tr(lang, "topic_diversity"), f"{tm.get(f'{m}_diversity', np.nan):.2f}")
    st.dataframe(topics_df, width="stretch", hide_index=True)
    weights = tm.get(f"{m}_doc_topic")
    if not (isinstance(weights, pd.DataFrame) and len(weights) == len(master)):
        return
    if selected_result_group and selected_result_group in master.columns:
        plot = pd.concat([master[[selected_result_group]].reset_index(drop=True), weights.reset_index(drop=True)], axis=1)
        plot = plot.groupby(selected_result_group, dropna=False, as_index=False)[list(weights.columns)].mean()
        long = plot.melt(id_vars=selected_result_group, var_name="topic", value_name="weight")
        st.plotly_chart(px.bar(long, x=selected_result_group, y="weight", color="topic",
                               title=f"{label} · {tr(lang,'topic_prevalence_group')}"), width="stretch", key=f"topics_group_{m}")
    else:
        plot = pd.concat([master[["doc_id"]].reset_index(drop=True), weights.reset_index(drop=True)], axis=1)
        long = plot.melt(id_vars="doc_id", var_name="topic", value_name="weight")
        st.plotly_chart(px.bar(long, x="doc_id", y="weight", color="topic", title=f"{label} · {tr(lang,'doc_topic_mixture')}"),
                        width="stretch", key=f"topics_doc_{m}")


with tab_topics:
    st.markdown(f"<div class='pm-note'>{tr(lang,'topic_note')}</div>", unsafe_allow_html=True)
    tm = R.get("topics", {})
    st.caption(tr(lang, "topic_unit_used") + ": **" + tr(lang, f"topic_unit_{R.get('topic_unit','document')}") + "**"
               + (f" · {len(R.get('topic_units', []))} " + tr(lang, "units") if R.get("topic_units") is not None else ""))
    c1, c2 = st.columns(2)
    with c1:
        render_topic_model("lda", tm)
    with c2:
        render_topic_model("nmf", tm)

    st.subheader(tr(lang, "k_diagnostics_title"))
    st.caption(tr(lang, "k_diagnostics_help"))
    d1, d2 = st.columns([3, 1])
    with d1:
        k_range = st.slider(tr(lang, "k_range"), 2, 30, (3, 15), key="pm_k_range")
    with d2:
        k_step = st.number_input(tr(lang, "k_step"), 1, 5, 2, key="pm_k_step")
    if st.button(tr(lang, "run_k_diagnostics"), key="pm_run_k_diag"):
        units = R.get("topic_units") or R["cleaned_docs"]
        ks = list(range(int(k_range[0]), int(k_range[1]) + 1, int(k_step)))
        with st.spinner(tr(lang, "running")):
            _kp = st.progress(0.0)
            st.session_state.pm_topic_k_diag = topic_k_diagnostics(
                units, ks, int(st.session_state.pm_params.get("max_features", 5000)),
                int(st.session_state.pm_params.get("random_state", 42)),
                progress=lambda f, msg: _kp.progress(min(1.0, f), text=msg))
            _kp.empty()
    kd = st.session_state.get("pm_topic_k_diag")
    if isinstance(kd, pd.DataFrame) and not kd.empty:
        c1, c2 = st.columns(2)
        with c1:
            st.plotly_chart(px.line(kd, x="k", y="mean_npmi", color="model", markers=True, title=tr(lang, "mean_npmi")), width="stretch", key="kd_npmi")
        with c2:
            st.plotly_chart(px.line(kd, x="k", y="topic_diversity", color="model", markers=True, title=tr(lang, "topic_diversity")), width="stretch", key="kd_div")
        st.dataframe(kd, width="stretch", hide_index=True)

with tab_sem:
    st.markdown(f"<div class='pm-note'>{tr(lang,'semantic_note')}</div>", unsafe_allow_html=True)
    if "semantic_error" in R:
        st.warning(f"{tr(lang,'model_unavailable')} ({R['semantic_error']})")
    elif "semantic" not in R:
        st.info(tr(lang,"model_unavailable"))
    else:
        sem = R["semantic"]
        k1,k2,k3 = st.columns(3)
        k1.metric("k", sem.get("k", 1))
        k2.metric(tr(lang,"silhouette"), "—" if not np.isfinite(sem.get("silhouette",np.nan)) else f"{sem['silhouette']:.3f}")
        k3.metric(tr(lang,"davies"), "—" if not np.isfinite(sem.get("davies_bouldin",np.nan)) else f"{sem['davies_bouldin']:.3f}")
        coords = sem["coords"]
        if R.get("semantic_unit") == "chunk":
            plot = chunks[["chunk_id","doc_id","language"]].copy()
            plot["x"], plot["y"], plot["cluster"] = coords[:,0], coords[:,1], sem["labels"].astype(str)
            hover = ["doc_id","language","chunk_id"]
        else:
            plot = master[["doc_id","language"]].copy()
            plot["x"], plot["y"], plot["cluster"] = coords[:,0], coords[:,1], sem["labels"].astype(str)
            hover = ["doc_id","language"]
        st.plotly_chart(px.scatter(plot, x="x", y="y", color="cluster", hover_data=hover, title=tr(lang,"semantic_clusters")), width="stretch")
        kd = sem.get("k_diagnostics")
        if isinstance(kd,pd.DataFrame) and not kd.empty:
            st.plotly_chart(px.line(kd, x="k", y="silhouette", markers=True, title="k diagnostics"), width="stretch")
        if "doc_similarity" in R and len(master) >= 2:
            st.subheader(tr(lang,"similarity"))
            sim = R["doc_similarity"]
            fig = go.Figure(data=go.Heatmap(z=sim.values, x=sim.columns, y=sim.index, zmin=-1, zmax=1, colorscale="RdBu", reversescale=True))
            fig.update_layout(height=max(450, min(900, 28*len(sim))))
            st.plotly_chart(fig, width="stretch")
            if R.get("hierarchy") is not None:
                try:
                    from scipy.cluster.hierarchy import dendrogram
                    d = dendrogram(R["hierarchy"], labels=master["doc_id"].astype(str).tolist(), no_plot=True)
                    xvals, yvals = [], []
                    for xs, ys in zip(d["icoord"], d["dcoord"]):
                        xvals.extend(xs + [None]); yvals.extend(ys + [None])
                    figd = go.Figure(go.Scatter(x=xvals, y=yvals, mode="lines", hoverinfo="skip"))
                    figd.update_layout(title="Hierarchical clustering · average linkage / cosine distance", xaxis=dict(tickmode="array", tickvals=[5+10*i for i in range(len(d["ivl"]))], ticktext=d["ivl"], tickangle=45), yaxis_title="Cosine distance", height=480)
                    st.plotly_chart(figd, width="stretch")
                except Exception:
                    pass
        st.subheader(tr(lang,"semantic_search"))
        q = st.text_input(tr(lang,"query"), key="semantic_query")
        if st.button(tr(lang,"search"), key="semantic_search_btn") and q.strip():
            try:
                emb = get_embedder().encode([q], normalize_embeddings=True)[0]
                scores = np.asarray(R["chunk_embeddings"]) @ emb
                ix = np.argsort(scores)[::-1][:10]
                out = chunks.iloc[ix][["chunk_id","doc_id","chunk_text"]].copy()
                out.insert(2, "similarity", scores[ix])
                st.dataframe(out, width="stretch", hide_index=True)
            except Exception as e:
                st.warning(str(e))

with tab_sent:
    st.markdown(f"<div class='pm-note'>{tr(lang,'sentiment_note')}</div>", unsafe_allow_html=True)
    if "sentiment_error" in R:
        st.warning(f"{tr(lang,'model_unavailable')} ({R['sentiment_error']})")
    elif "sentiment_docs" not in R:
        st.info(tr(lang,"model_unavailable"))
    else:
        sd = R["sentiment_docs"]
        st.subheader(tr(lang,"sentiment_doc"))
        st.dataframe(sd, width="stretch", hide_index=True)
        if selected_result_group and selected_result_group in R["docs"].columns:
            sgroup = sd.merge(R["docs"][["doc_id",selected_result_group]], on="doc_id", how="left")
            ssum = grouped_numeric_summary(sgroup, selected_result_group, "sentiment_score", n_boot=1500, random_state=int(st.session_state.pm_params.get("random_state",42)))
            if not ssum.empty:
                ssum["err_plus"] = ssum["ci_high_95"] - ssum["mean"]
                ssum["err_minus"] = ssum["mean"] - ssum["ci_low_95"]
                st.plotly_chart(px.bar(ssum, x=selected_result_group, y="mean", error_y="err_plus", error_y_minus="err_minus", range_y=[-1,1], title=tr(lang,"sentiment_by_group")), width="stretch")
                st.plotly_chart(px.box(sgroup, x=selected_result_group, y="sentiment_score", points="all", range_y=[-1,1]), width="stretch")
        else:
            st.plotly_chart(px.bar(sd, x="doc_id", y="sentiment_score", color="sentiment", range_y=[-1,1]), width="stretch")
        st.subheader(tr(lang,"sentiment_chunks"))
        sc = R["sentiment_chunks"].merge(chunks[["chunk_id","language"]], on="chunk_id", how="left")
        st.plotly_chart(px.histogram(sc, x="sentiment_score", color="sentiment", nbins=30, marginal="box"), width="stretch")

with tab_frames:
    st.markdown(f"<div class='pm-note'>{tr(lang,'frame_note')}</div>", unsafe_allow_html=True)
    raw = st.text_area(tr(lang,"frame_dictionary"), value="", height=220,
                       placeholder=tr(lang,"frame_dictionary_placeholder"), key="pm_dictionary_raw")
    st.caption(tr(lang, "frame_dictionary_syntax"))
    dict_accents = st.checkbox(tr(lang, "accent_insensitive"), value=True, key="pm_dict_accents")
    try:
        fdict = {} if not raw.strip() else json.loads(raw)
        if not isinstance(fdict, dict):
            raise ValueError(tr(lang,"frame_dictionary_must_object"))
        fdict = {str(k): ([v] if isinstance(v, str) else list(v)) for k, v in fdict.items()}
        if not fdict:
            st.session_state.pop("pm_custom_frames", None)
            st.session_state.pop("pm_custom_frames_audit", None)
            st.info(tr(lang,"frame_dictionary_empty"))
        else:
            fdf = frame_indicators(R["docs"], fdict, accent_insensitive=dict_accents)
            st.session_state.pm_custom_frames = fdf
            audit = dictionary_term_counts(R["docs"], fdict, accent_insensitive=dict_accents)
            st.session_state.pm_custom_frames_audit = audit
            rate_cols = [c for c in fdf.columns if c.endswith("_per_1000")]
            st.dataframe(fdf, width="stretch", hide_index=True)
            if rate_cols:
                if selected_result_group and selected_result_group in R["docs"].columns:
                    fg = fdf.merge(R["docs"][["doc_id",selected_result_group]], on="doc_id", how="left")
                    gmeans = fg.groupby(selected_result_group, dropna=False, as_index=False)[rate_cols].mean()
                    melt = gmeans.melt(id_vars=selected_result_group, value_vars=rate_cols, var_name="frame", value_name="per_1000")
                    st.plotly_chart(px.bar(melt, x=selected_result_group, y="per_1000", color="frame", barmode="group", title=tr(lang,"dictionary_by_group")), width="stretch")
                else:
                    melt = fdf.melt(id_vars="doc_id", value_vars=rate_cols, var_name="frame", value_name="per_1000")
                    st.plotly_chart(px.bar(melt, x="doc_id", y="per_1000", color="frame", barmode="group"), width="stretch")
            with st.expander(tr(lang, "dictionary_audit")):
                st.caption(tr(lang, "dictionary_audit_help"))
                if audit.empty:
                    st.info(tr(lang, "not_enough"))
                else:
                    by_entry = audit.groupby(["category", "entry"], as_index=False).agg(hits=("hits", "sum"), documents=("doc_id", "nunique"))
                    st.dataframe(by_entry.sort_values("hits", ascending=False), width="stretch", hide_index=True)
                    st.dataframe(audit, width="stretch", hide_index=True)
    except Exception as e:
        st.error(str(e))
    st.subheader(tr(lang,"concordance"))
    c1,c2 = st.columns([3,1])
    with c1: term = st.text_input(tr(lang,"term"), key="kwic_term", help=tr(lang, "kwic_help"))
    with c2: ctx = st.number_input(tr(lang,"context_words"), min_value=5, max_value=40, value=12)
    if term.strip():
        st.dataframe(kwic(R["docs"], term, int(ctx), 300, accent_insensitive=dict_accents), width="stretch", hide_index=True)


with tab_concepts:
    st.markdown(f"<div class='pm-note'>{tr(lang,'concepts_note')}</div>", unsafe_allow_html=True)
    if "chunk_embeddings" not in R:
        st.info(tr(lang, "supervised_need_semantic"))
    else:
        anchors_raw = st.text_area(tr(lang, "concept_anchors"), value="", height=220,
                                   placeholder=tr(lang, "concept_anchors_placeholder"), key="pm_concepts_raw")
        st.caption(tr(lang, "concept_anchors_syntax"))
        c1, c2 = st.columns(2)
        with c1:
            c_thr = st.slider(tr(lang, "concept_threshold"), 0.10, 0.90, 0.45, 0.01, key="pm_concept_thr",
                              help=tr(lang, "concept_threshold_help"))
        with c2:
            c_top = st.slider(tr(lang, "concept_top_passages"), 3, 30, 8, key="pm_concept_top")
        if st.button(tr(lang, "run_concepts"), type="primary", key="pm_run_concepts"):
            try:
                anchors = parse_concept_anchors(anchors_raw)
                if len(anchors) < 1:
                    st.warning(tr(lang, "concept_need_anchors"))
                else:
                    with st.spinner(tr(lang, "running")):
                        embedder = get_embedder()
                        cemb = {k: embedder.encode(v, normalize_embeddings=True, show_progress_bar=False) for k, v in anchors.items()}
                        out = concept_projection(np.asarray(R["chunk_embeddings"]), chunks, R["docs"]["doc_id"].astype(str).tolist(),
                                                 cemb, threshold=float(c_thr), top_passages=int(c_top))
                    st.session_state.pm_concepts = out
                    st.session_state.pm_concepts_docs = out["documents"]
                    st.session_state.pm_concepts_info = pd.DataFrame(
                        [{"concept": k, "anchors": " | ".join(v), "threshold": float(c_thr), "embedding_model": EMBEDDING_MODEL}
                         for k, v in anchors.items()])
            except Exception as e:
                st.error(str(e))
        out = st.session_state.get("pm_concepts")
        if isinstance(out, dict) and not out["documents"].empty:
            cdocs = out["documents"]
            names = [c[len("concept_"):-len("_mean")] for c in cdocs.columns if c.startswith("concept_") and c.endswith("_mean")]
            measure = st.radio(tr(lang, "concept_measure"), ["centred", "mean", "share"], horizontal=True,
                               format_func=lambda x: tr(lang, f"concept_measure_{x}"), key="pm_concept_measure")
            cols = [f"concept_{n}_{measure}" for n in names]
            view = cdocs.merge(R["docs"][["doc_id"] + ([selected_result_group] if selected_result_group else [])].astype({"doc_id": str}),
                               on="doc_id", how="left")
            if selected_result_group:
                gm = view.groupby(selected_result_group, dropna=False)[cols].mean()
                gm.columns = names
                fig = go.Figure(go.Heatmap(z=gm.values, x=names, y=gm.index.astype(str),
                                           colorscale="RdBu" if measure == "centred" else "Viridis",
                                           zmid=0 if measure == "centred" else None))
                fig.update_layout(title=tr(lang, "concepts_by_group"), height=max(380, 45 * len(gm)), yaxis_title=selected_result_group)
                st.plotly_chart(fig, width="stretch", key="concept_heat_group")
            else:
                hm = view.set_index("doc_id")[cols]
                hm.columns = names
                fig = go.Figure(go.Heatmap(z=hm.values, x=names, y=hm.index.astype(str),
                                           colorscale="RdBu" if measure == "centred" else "Viridis",
                                           zmid=0 if measure == "centred" else None))
                fig.update_layout(title=tr(lang, "concepts_by_document"), height=max(380, 22 * len(hm)))
                st.plotly_chart(fig, width="stretch", key="concept_heat_doc")
            st.dataframe(cdocs, width="stretch", hide_index=True)
            st.subheader(tr(lang, "concept_passages"))
            st.caption(tr(lang, "concept_passages_help"))
            sel = st.selectbox(tr(lang, "concept"), names, key="pm_concept_sel")
            st.dataframe(out["passages"].query("concept == @sel"), width="stretch", hide_index=True)


with tab_entities:
    st.markdown(f"<div class='pm-note'>{tr(lang,'entities_note')}</div>", unsafe_allow_html=True)
    if "entities_error" in R:
        st.warning(f"{tr(lang,'model_unavailable')} ({R['entities_error']})")
    elif "entities" not in R:
        st.info(tr(lang,"entities_not_run"))
    else:
        entities = R["entities"]
        if entities.empty:
            st.info(tr(lang,"not_enough"))
        else:
            types = sorted(entities["entity_type"].dropna().astype(str).unique().tolist())
            selected_types = st.multiselect(tr(lang,"entity_types"), types, default=[x for x in types if x in {"PER","ORG","LOC"}] or types)
            ef = entities[entities["entity_type"].isin(selected_types)] if selected_types else entities.iloc[0:0]
            totals = ef.groupby(["entity","entity_type"], as_index=False).agg(mentions=("mentions","sum"), documents=("doc_id","nunique"), mean_score=("mean_score","mean")).sort_values(["mentions","documents"], ascending=False)
            c1,c2=st.columns([1,2])
            with c1:
                st.subheader(tr(lang,"top_entities")); st.dataframe(totals.head(50), width="stretch",hide_index=True)
            with c2:
                if not totals.empty:
                    st.plotly_chart(px.bar(totals.head(25).sort_values("mentions"),x="mentions",y="entity",color="entity_type",orientation="h"), width="stretch")
            st.subheader(tr(lang,"actor_topic_network"))
            tm=R.get("topics",{})
            topic_options=[]
            if isinstance(tm.get("nmf_doc_topic"),pd.DataFrame): topic_options.append("NMF")
            if isinstance(tm.get("lda_doc_topic"),pd.DataFrame): topic_options.append("LDA")
            if topic_options and not ef.empty:
                c1,c2,c3=st.columns(3)
                with c1: topic_method=st.selectbox(tr(lang,"topic_model"),topic_options,key="network_topic_model")
                with c2: top_e=st.slider(tr(lang,"network_top_entities"),5,60,25,5)
                with c3: min_w=st.number_input(tr(lang,"minimum_edge"),min_value=0.0,value=0.10,step=0.05)
                tdf=tm["nmf_doc_topic"] if topic_method=="NMF" else tm["lda_doc_topic"]
                edges=actor_topic_edges(ef,tdf,R["docs"]["doc_id"].astype(str).tolist(),selected_types,top_e,float(min_w))
                if edges.empty:
                    st.info(tr(lang,"not_enough"))
                else:
                    fig=actor_topic_network_figure(edges)
                    if fig is not None: st.plotly_chart(fig, width="stretch")
                    with st.expander(tr(lang,"network_edges")):
                        st.dataframe(edges, width="stretch",hide_index=True)
                    st.session_state.pm_network_edges=edges
            else:
                st.info(tr(lang,"not_enough"))


with tab_grouped:
    st.markdown(f"<div class='pm-note'>{tr(lang,'grouped_note')}</div>", unsafe_allow_html=True)
    if not selected_result_group or selected_result_group not in R["docs"].columns:
        st.info(tr(lang, "choose_result_group"))
    else:
        gcol = selected_result_group
        group_counts = R["docs"].groupby(gcol, dropna=False, as_index=False).agg(documents=("doc_id","nunique"))
        words_by_doc = R["doc_stats"][["doc_id","words"]]
        group_counts = group_counts.merge(
            R["docs"][["doc_id",gcol]].merge(words_by_doc,on="doc_id",how="left").groupby(gcol,dropna=False,as_index=False)["words"].sum(),
            on=gcol, how="left"
        )
        c1,c2 = st.columns([1,2])
        with c1:
            st.subheader(tr(lang,"group_composition"))
            st.dataframe(group_counts, width="stretch", hide_index=True)
        with c2:
            st.plotly_chart(px.bar(group_counts, x=gcol, y="documents", hover_data=["words"], title=tr(lang,"documents_by_group")), width="stretch")

        analysis_df = build_analysis_dataset(R)
        numeric_outcomes = [
            c for c in analysis_df.columns
            if c not in {"doc_id","text"} and pd.api.types.is_numeric_dtype(analysis_df[c])
            and analysis_df[c].notna().sum() >= 2 and analysis_df[c].nunique(dropna=True) > 1
        ]
        if numeric_outcomes:
            st.subheader(tr(lang,"grouped_numeric_results"))
            default_idx = numeric_outcomes.index("sentiment_score") if "sentiment_score" in numeric_outcomes else 0
            outcome = st.selectbox(tr(lang,"outcome"), numeric_outcomes, index=default_idx, key="pm_grouped_outcome")
            summary = grouped_numeric_summary(
                analysis_df, gcol, outcome, n_boot=1500,
                random_state=int(st.session_state.pm_params.get("random_state",42))
            )
            st.session_state.pm_grouped_summary = summary.assign(outcome=str(outcome), group_variable=str(gcol)) if not summary.empty else summary
            if not summary.empty:
                summary_plot = summary.copy()
                summary_plot["err_plus"] = summary_plot["ci_high_95"] - summary_plot["mean"]
                summary_plot["err_minus"] = summary_plot["mean"] - summary_plot["ci_low_95"]
                c1,c2 = st.columns([1,2])
                with c1:
                    st.dataframe(summary.drop(columns=[], errors="ignore"), width="stretch", hide_index=True)
                with c2:
                    st.plotly_chart(px.bar(
                        summary_plot, x=gcol, y="mean", error_y="err_plus", error_y_minus="err_minus",
                        title=f"{outcome} · {tr(lang,'group_mean_ci')}"
                    ), width="stretch")
                box_df = analysis_df[[gcol,outcome]].dropna().copy()
                if not box_df.empty:
                    st.plotly_chart(px.box(box_df, x=gcol, y=outcome, points="all", title=tr(lang,"group_distribution")), width="stretch")

        tm = R.get("topics", {})
        topic_method_options = []
        if isinstance(tm.get("nmf_doc_topic"), pd.DataFrame): topic_method_options.append("NMF")
        if isinstance(tm.get("lda_doc_topic"), pd.DataFrame): topic_method_options.append("LDA")
        if topic_method_options:
            st.subheader(tr(lang,"topics_by_group"))
            topic_method = st.radio(tr(lang,"topic_model"), topic_method_options, horizontal=True, key="pm_grouped_topic_model")
            tdf = tm["nmf_doc_topic"] if topic_method == "NMF" else tm["lda_doc_topic"]
            if len(tdf) == len(R["docs"]):
                topic_group = pd.concat([R["docs"][[gcol]].reset_index(drop=True), tdf.reset_index(drop=True)], axis=1)
                topic_means = topic_group.groupby(gcol, dropna=False).mean(numeric_only=True)
                st.dataframe(topic_means.reset_index(), width="stretch", hide_index=True)
                if not topic_means.empty:
                    fig = go.Figure(go.Heatmap(z=topic_means.values, x=topic_means.columns, y=topic_means.index.astype(str), colorscale="Viridis"))
                    fig.update_layout(title=tr(lang,"topic_prevalence_group"), xaxis_title=tr(lang,"topics_tab"), yaxis_title=gcol, height=max(380, 55*len(topic_means)))
                    st.plotly_chart(fig, width="stretch")

        vals = [x for x in R["docs"][gcol].dropna().unique().tolist()]
        if len(vals) >= 2:
            st.subheader(tr(lang,"distinctive_language_group"))
            focus_group = st.selectbox(tr(lang,"focus_group"), vals, key="pm_focus_group")
            ids_focus = set(R["docs"].loc[R["docs"][gcol] == focus_group, "doc_id"].astype(str))
            cleaned_map = dict(zip(R["docs"]["doc_id"].astype(str), R["cleaned_docs"]))
            a = [cleaned_map[i] for i in ids_focus if i in cleaned_map]
            b = [cleaned_map[i] for i in cleaned_map if i not in ids_focus]
            if a and b:
                key_group = informative_log_odds(a, b, 40)
                st.caption(tr(lang,"focus_group_direction"))
                st.dataframe(key_group, width="stretch", hide_index=True)

with tab_groups:
    st.markdown(f"<div class='pm-note'>{tr(lang,'group_note')}</div>", unsafe_allow_html=True)
    analysis_df = build_analysis_dataset(R)
    excluded = {"doc_id","text","language","sentiment","sentiment_score","p_positive","p_neutral","p_negative","sentiment_sd","sentiment_chunks"}
    group_cols = [c for c in R["docs"].columns if c not in excluded and R["docs"][c].nunique(dropna=True) >= 2 and R["docs"][c].nunique(dropna=True) <= 30]
    if not group_cols:
        st.info(tr(lang,"not_enough"))
    else:
        gcol = st.selectbox(tr(lang,"group_column"), group_cols)
        vals = [x for x in R["docs"][gcol].dropna().unique().tolist()]
        c1,c2 = st.columns(2)
        with c1: ga = st.selectbox(tr(lang,"group_a"), vals, index=0)
        with c2: gb = st.selectbox(tr(lang,"group_b"), vals, index=1 if len(vals)>1 else 0)
        if ga == gb:
            st.warning(tr(lang,"different_groups"))
        else:
            ids_a = set(R["docs"].loc[R["docs"][gcol] == ga, "doc_id"].astype(str))
            ids_b = set(R["docs"].loc[R["docs"][gcol] == gb, "doc_id"].astype(str))
            cleaned_map = dict(zip(R["docs"]["doc_id"].astype(str), R["cleaned_docs"]))
            keydf = informative_log_odds([cleaned_map[i] for i in ids_a], [cleaned_map[i] for i in ids_b], 50)
            st.subheader(tr(lang,"keyness")); st.caption(tr(lang,"keyness_direction")); st.dataframe(keydf, width="stretch",hide_index=True)
            numeric_outcomes=[c for c in analysis_df.columns if c not in {"doc_id","text"} and pd.api.types.is_numeric_dtype(analysis_df[c]) and analysis_df[c].notna().sum()>=4 and analysis_df[c].nunique(dropna=True)>1]
            if numeric_outcomes:
                default_idx=numeric_outcomes.index("sentiment_score") if "sentiment_score" in numeric_outcomes else 0
                outcome=st.selectbox(tr(lang,"outcome"),numeric_outcomes,index=default_idx,key="group_numeric_outcome")
                tmp=analysis_df[["doc_id",gcol,outcome]].copy()
                a=pd.to_numeric(tmp.loc[tmp[gcol]==ga,outcome],errors="coerce").dropna().values
                b=pd.to_numeric(tmp.loc[tmp[gcol]==gb,outcome],errors="coerce").dropna().values
                testdf=numeric_group_test(a,b); bootdf=bootstrap_difference_ci(a,b,n_boot=3000,random_state=int(st.session_state.pm_params.get("random_state",42)))
                c1,c2=st.columns(2)
                with c1:
                    st.subheader(tr(lang,"inferential_tests")); st.dataframe(testdf, width="stretch",hide_index=True) if not testdf.empty else st.info(tr(lang,"not_enough"))
                with c2:
                    st.subheader(tr(lang,"bootstrap_ci")); st.dataframe(bootdf, width="stretch",hide_index=True) if not bootdf.empty else st.info(tr(lang,"not_enough"))
                if not bootdf.empty:
                    st.session_state.pm_group_bootstrap=bootdf.assign(group_variable=str(gcol),group_A=str(ga),group_B=str(gb),outcome=str(outcome))


with tab_temporal:
    st.markdown(f"<div class='pm-note'>{tr(lang,'temporal_note')}</div>", unsafe_allow_html=True)
    analysis_df=build_analysis_dataset(R)
    meta_cols=[c for c in R["docs"].columns if c not in {"doc_id","text","source_file","language","sentiment","sentiment_score","p_positive","p_neutral","p_negative","sentiment_sd","sentiment_chunks"}]
    time_candidates=[]
    for c in meta_cols:
        name=str(c).lower()
        ser=R["docs"][c]
        numeric=pd.to_numeric(ser,errors="coerce")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            parsed=pd.to_datetime(ser,errors="coerce") if not pd.api.types.is_numeric_dtype(ser) else pd.Series(pd.NaT,index=ser.index)
        if any(k in name for k in ["year","date","time","wave","month","año","fecha","anno","data"]) or numeric.notna().mean()>=.8 or parsed.notna().mean()>=.8:
            if ser.nunique(dropna=True)>=2: time_candidates.append(c)
    outcomes=[c for c in analysis_df.columns if pd.api.types.is_numeric_dtype(analysis_df[c]) and analysis_df[c].notna().sum()>=4 and analysis_df[c].nunique(dropna=True)>1]
    if not time_candidates or not outcomes:
        st.info(tr(lang,"temporal_need"))
    else:
        c1,c2,c3=st.columns(3)
        with c1: tcol=st.selectbox(tr(lang,"time_variable"),time_candidates)
        with c2:
            t_outcomes=[c for c in outcomes if c!=tcol]
            default_idx=t_outcomes.index("sentiment_score") if "sentiment_score" in t_outcomes else 0
            ycol=st.selectbox(tr(lang,"outcome"),t_outcomes,index=default_idx,key="time_outcome")
        with c3:
            gopts=["—"]+[c for c in meta_cols if c!=tcol and R["docs"][c].nunique(dropna=True) in range(2,16)]
            tg=st.selectbox(tr(lang,"temporal_group"),gopts)
        group=None if tg=="—" else tg
        summary=temporal_bootstrap_summary(analysis_df,tcol,ycol,group_col=group,n_boot=1500,random_state=int(st.session_state.pm_params.get("random_state",42)))
        trend=robust_time_trend(analysis_df,tcol,ycol,group_col=group)
        if summary.empty:
            st.info(tr(lang,"not_enough"))
        else:
            summary["err_plus"]=summary["ci_high"]-summary["mean"]; summary["err_minus"]=summary["mean"]-summary["ci_low"]
            fig=px.line(summary,x="period",y="mean",color=group if group else None,markers=True,error_y="err_plus",error_y_minus="err_minus",title=tr(lang,"temporal_evolution"))
            st.plotly_chart(fig, width="stretch")
            st.subheader(tr(lang,"robust_trend")); st.dataframe(trend, width="stretch",hide_index=True) if not trend.empty else st.info(tr(lang,"not_enough"))
            with st.expander(tr(lang,"bootstrap_summary")): st.dataframe(summary.drop(columns=["err_plus","err_minus"]), width="stretch",hide_index=True)
            st.session_state.pm_temporal_summary=summary.drop(columns=["err_plus","err_minus"])
            st.session_state.pm_temporal_trend=trend

with tab_supervised:
    st.markdown(f"<div class='pm-note'>{tr(lang,'supervised_note')}</div>", unsafe_allow_html=True)
    if "doc_embeddings" not in R:
        st.info(tr(lang,"supervised_need_semantic"))
    else:
        label_candidates=[]
        for c in R["docs"].columns:
            if c in {"doc_id","text","source_file","language","sentiment","sentiment_score","p_positive","p_neutral","p_negative","sentiment_sd","sentiment_chunks"}: continue
            n=R["docs"][c].nunique(dropna=True)
            counts=R["docs"][c].dropna().astype(str).value_counts()
            if 2<=n<=12 and len(counts)>=2: label_candidates.append(c)
        if not label_candidates:
            st.info(tr(lang,"supervised_need_labels"))
        else:
            c1,c2=st.columns(2)
            with c1: task=st.selectbox(tr(lang,"supervised_task"),[tr(lang,"stance_task"),tr(lang,"framing_task"),tr(lang,"other_task")])
            with c2: label_col=st.selectbox(tr(lang,"label_column"),label_candidates)
            target=st.text_input(tr(lang,"stance_target"),placeholder=tr(lang,"stance_target_placeholder"))
            if st.button(tr(lang,"train_classifier"),type="primary",key="train_supervised"):
                out=supervised_embedding_classifier(np.asarray(R["doc_embeddings"]),R["docs"]["doc_id"].astype(str).tolist(),R["docs"][label_col].tolist(),random_state=int(st.session_state.pm_params.get("random_state",42)))
                if "error" in out:
                    st.warning(out["error"])
                else:
                    st.session_state.pm_supervised=out
                    st.session_state.pm_supervised_info=pd.DataFrame([{"task":task,"label_column":label_col,"target":target,"embedding_model":EMBEDDING_MODEL}])
            out=st.session_state.get("pm_supervised")
            if isinstance(out,dict) and "metrics" in out:
                st.subheader(tr(lang,"cv_performance")); st.dataframe(out["metrics"], width="stretch",hide_index=True)
                c1,c2=st.columns(2)
                with c1:
                    st.subheader(tr(lang,"confusion_matrix")); cm=out["confusion"].set_index("actual"); st.plotly_chart(px.imshow(cm,text_auto=True,aspect="auto",labels=dict(x="Predicted",y="Actual",color="N")), width="stretch")
                with c2:
                    st.subheader(tr(lang,"classification_report")); st.dataframe(out["classification_report"], width="stretch",hide_index=True)
                st.subheader(tr(lang,"predictions")); st.dataframe(out["predictions"], width="stretch",hide_index=True)

with tab_models:
    st.markdown(f"<div class='pm-note'>{tr(lang,'model_note')}</div>", unsafe_allow_html=True)
    # Build a document-level modelling dataset from metadata + generated indicators.
    model_df = build_analysis_dataset(R)

    outcome_candidates = []
    for c in model_df.columns:
        if c in {"text", "doc_id", "source_file"}: continue
        if pd.api.types.is_numeric_dtype(model_df[c]) and model_df[c].notna().sum() >= 4 and model_df[c].nunique(dropna=True) > 1:
            if (c.startswith("frame_") and c.endswith("_per_1000")) or c.startswith("concept_") \
                    or c in {"sentiment_score","mattr_50","type_token_ratio","avg_sentence_words"} or c.startswith("LDA_") or c.startswith("NMF_"):
                outcome_candidates.append(c)
    if not outcome_candidates:
        st.info(tr(lang,"not_enough"))
    else:
        ycol = st.selectbox(tr(lang,"outcome"), outcome_candidates)
        predictor_candidates = []
        for c in R["docs"].columns:
            if c in {"doc_id","text","sentiment","sentiment_score","p_positive","p_neutral","p_negative","sentiment_sd","sentiment_chunks"}: continue
            nun = R["docs"][c].nunique(dropna=True)
            if 1 < nun <= max(30, len(R["docs"])//2) or pd.api.types.is_numeric_dtype(R["docs"][c]):
                predictor_candidates.append(c)
        xcols = st.multiselect(tr(lang,"predictors"), predictor_candidates)
        if xcols and st.button(tr(lang,"fit_model"), key="fit_ols"):
            try:
                import statsmodels.api as sm
                raw = model_df[[ycol] + xcols].copy()
                # Treat low-cardinality nonnumeric variables as categorical; numeric variables remain continuous.
                X = pd.get_dummies(raw[xcols], columns=[c for c in xcols if not pd.api.types.is_numeric_dtype(raw[c])], drop_first=True, dtype=float)
                X = X.apply(pd.to_numeric, errors="coerce")
                y = pd.to_numeric(raw[ycol], errors="coerce")
                dat = pd.concat([y.rename(ycol), X], axis=1).replace([np.inf,-np.inf], np.nan).dropna()
                X2 = sm.add_constant(dat.drop(columns=[ycol]), has_constant="add")
                if len(dat) <= X2.shape[1] + 2:
                    st.warning(tr(lang,"not_enough"))
                else:
                    fit = sm.OLS(dat[ycol], X2).fit(cov_type="HC3")
                    ci = fit.conf_int()
                    coef = pd.DataFrame({
                        "term": fit.params.index,
                        "coefficient": fit.params.values,
                        "robust_se_HC3": fit.bse.values,
                        "t": fit.tvalues.values,
                        "p_value": fit.pvalues.values,
                        "ci_low_95": ci[0].values,
                        "ci_high_95": ci[1].values,
                    })
                    m1,m2,m3 = st.columns(3)
                    m1.metric("N", int(fit.nobs)); m2.metric("R²", f"{fit.rsquared:.3f}"); m3.metric("Adj. R²", f"{fit.rsquared_adj:.3f}")
                    st.subheader(tr(lang,"coef_table"))
                    st.dataframe(coef, width="stretch", hide_index=True)
                    st.session_state.pm_model_coefficients = coef
                    st.session_state.pm_model_info = pd.DataFrame([{"outcome":ycol,"predictors":", ".join(xcols),"n":int(fit.nobs),"r_squared":fit.rsquared,"adj_r_squared":fit.rsquared_adj,"covariance":"HC3"}])
            except Exception as e:
                st.error(str(e))


with tab_publication:
    st.markdown(f"<div class='pm-note'>{tr(lang,'publication_note')}</div>", unsafe_allow_html=True)
    html_report=make_html_report(R,st.session_state.pm_params,lang)
    c1,c2=st.columns(2)
    with c1:
        st.download_button(tr(lang,"download_report"),html_report,"PerspectiveMapper_v3_3_report.html","text/html", width="stretch")
    with c2:
        if st.button(tr(lang,"prepare_publication_package"), width="stretch",key="prepare_pub_package"):
            with st.spinner(tr(lang,"preparing_package")):
                st.session_state.pm_publication_package=publication_package(R,st.session_state.pm_params)
        if "pm_publication_package" in st.session_state:
            st.download_button(tr(lang,"download_publication_package"),st.session_state.pm_publication_package,"PerspectiveMapper_v3_3_publication_package.zip","application/zip", width="stretch")
    st.caption(tr(lang,"publication_contents"))

with tab_export:
    meth = methodology_table(st.session_state.pm_params, R)
    st.subheader(tr(lang,"methodology"))
    st.dataframe(meth, width="stretch", hide_index=True)
    st.subheader(tr(lang,"limitations"))
    st.markdown("""
- Topic labels are inferred from high-weight terms and still require substantive interpretation by the researcher.
- Multilingual lexical models operate on surface forms; cross-language equivalence is provided primarily by the multilingual embedding model.
- Sentiment models can be domain-sensitive and may perform differently across languages, registers and demographic groups.
- Framing dictionaries measure lexical prevalence, not causality, intention, fairness or discriminatory bias.
- Significance tests assume suitable independent document-level units; repeated measures or nested corpora require an appropriate multilevel design outside this dashboard.
    """)
    sheets = {
        "Documents": master.drop(columns=["text"], errors="ignore"),
        "Document_stats": R["doc_stats"],
        "Chunks": chunks,
        "Keywords_TFIDF": R["keywords"],
        "Collocations_PMI": R["collocations"],
        "Frames": R["frames"],
        "Methodology": meth,
    }
    tm = R.get("topics", {})
    for key, name in [("lda_topics","LDA_topics"),("lda_doc_topic","LDA_doc_topic"),("nmf_topics","NMF_topics"),("nmf_doc_topic","NMF_doc_topic")]:
        if key in tm: sheets[name] = tm[key]
    if "sentiment_docs" in R: sheets["Sentiment_docs"] = R["sentiment_docs"]
    if "sentiment_chunks" in R: sheets["Sentiment_chunks"] = R["sentiment_chunks"]
    if "entity_mentions" in R: sheets["Entity_mentions"] = R["entity_mentions"]
    if "entities" in R: sheets["Entities"] = R["entities"]
    if "keyword_filter_audit" in R and isinstance(R["keyword_filter_audit"], pd.DataFrame): sheets["Keyword_filter_audit"] = R["keyword_filter_audit"]
    if "pm_custom_frames" in st.session_state and isinstance(st.session_state.pm_custom_frames, pd.DataFrame): sheets["Custom_dictionary"] = st.session_state.pm_custom_frames
    if "pm_grouped_summary" in st.session_state and isinstance(st.session_state.pm_grouped_summary, pd.DataFrame): sheets["Grouped_summary"] = st.session_state.pm_grouped_summary
    if "pm_group_bootstrap" in st.session_state: sheets["Group_bootstrap"] = st.session_state.pm_group_bootstrap
    if "pm_temporal_summary" in st.session_state: sheets["Temporal_summary"] = st.session_state.pm_temporal_summary
    if "pm_temporal_trend" in st.session_state: sheets["Temporal_trend"] = st.session_state.pm_temporal_trend
    if "pm_network_edges" in st.session_state: sheets["Actor_topic_edges"] = st.session_state.pm_network_edges
    if "pm_supervised" in st.session_state and isinstance(st.session_state.pm_supervised,dict):
        for _k,_name in [("metrics","Supervised_metrics"),("confusion","Supervised_confusion"),("classification_report","Supervised_report"),("predictions","Supervised_predictions")]:
            if isinstance(st.session_state.pm_supervised.get(_k),pd.DataFrame): sheets[_name]=st.session_state.pm_supervised[_k]
    if "pm_supervised_info" in st.session_state: sheets["Supervised_info"] = st.session_state.pm_supervised_info
    if "doc_similarity" in R:
        sim_export = R["doc_similarity"].reset_index().rename(columns={"doc_id":"doc_id"})
        sheets["Similarity"] = sim_export
    if "semantic" in R and isinstance(R["semantic"].get("k_diagnostics"), pd.DataFrame):
        sheets["Cluster_diagnostics"] = R["semantic"]["k_diagnostics"]
    if "pm_model_coefficients" in st.session_state:
        sheets["OLS_coefficients"] = st.session_state.pm_model_coefficients
    if "pm_model_info" in st.session_state:
        sheets["OLS_model_info"] = st.session_state.pm_model_info
    if "pm_custom_frames_audit" in st.session_state and isinstance(st.session_state.pm_custom_frames_audit, pd.DataFrame):
        sheets["Dictionary_audit"] = st.session_state.pm_custom_frames_audit
    if isinstance(st.session_state.get("pm_concepts"), dict):
        sheets["Concepts_documents"] = st.session_state.pm_concepts["documents"]
        sheets["Concepts_passages"] = st.session_state.pm_concepts["passages"]
        sheets["Concepts_chunks"] = st.session_state.pm_concepts["chunks"]
    if "pm_concepts_info" in st.session_state: sheets["Concepts_anchors"] = st.session_state.pm_concepts_info
    if "pm_topic_k_diag" in st.session_state: sheets["Topic_k_diagnostics"] = st.session_state.pm_topic_k_diag
    for _m in ["lda", "nmf"]:
        if isinstance(tm.get(f"{_m}_unit_topic"), pd.DataFrame): sheets[f"{_m.upper()}_passage_topic"] = tm[f"{_m}_unit_topic"]
    c1,c2,c3 = st.columns(3)
    with c1:
        # Building the workbook can take a while for large corpora, so it is created on demand
        # rather than on every rerun.
        if st.button(tr(lang, "prepare_excel"), width="stretch", key="pm_prepare_excel"):
            with st.spinner(tr(lang, "running")):
                st.session_state.pm_excel = excel_bytes(sheets)
        if "pm_excel" in st.session_state:
            st.download_button(tr(lang,"download_excel"), st.session_state.pm_excel, "PerspectiveMapper_v3_3_results.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
    with c2:
        csv = master.drop(columns=["text"], errors="ignore").to_csv(index=False).encode("utf-8-sig")
        st.download_button(tr(lang,"download_csv"), csv, "PerspectiveMapper_documents.csv", "text/csv", width="stretch")
    with c3:
        payload = {
            "parameters": st.session_state.pm_params,
            "documents": json_safe(master.drop(columns=["text"], errors="ignore")),
            "topics": json_safe(R.get("topics", {})),
            "keywords": json_safe(R.get("keywords")),
            "keyword_filter_audit": json_safe(R.get("keyword_filter_audit")),
            "grouped_summary": json_safe(st.session_state.get("pm_grouped_summary")),
            "entities": json_safe(R.get("entities")),
            "supervised": json_safe(st.session_state.get("pm_supervised")),
            "temporal_summary": json_safe(st.session_state.get("pm_temporal_summary")),
            "topic_unit": R.get("topic_unit"),
            "topic_k_diagnostics": json_safe(st.session_state.get("pm_topic_k_diag")),
            "concepts": json_safe(st.session_state.get("pm_concepts_docs")),
            "concept_anchors": json_safe(st.session_state.get("pm_concepts_info")),
        }
        st.download_button(tr(lang,"download_json"), json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8"), "PerspectiveMapper_v3_3_results.json", "application/json", width="stretch")

st.markdown("---")
st.caption("PerspectiveMapper v3.3 Research Edition · multilingual, inferential and reproducible text analysis")
